import logging

import openpyxl


_log = logging.getLogger(__name__)


def iter_rows(ws):
    indexed_rows = enumerate(ws.iter_rows(values_only=True), start=1)
    for i, row in indexed_rows:
        cell = row[0]
        if cell == 'Candidate':
            break

    i, row = next(indexed_rows)
    cell = row[0]
    assert cell is None

    is_candidate = True
    for i, row in indexed_rows:
        cell = row[0]
        if cell is None:
            break
        if cell == 'Continuing Ballots Total':
            is_candidate = False

        yield (i, cell, is_candidate)


def parse_sheet_2(wb):
    ws = wb['Sheet2']

    candidates = []
    subtotals = []
    non_candidate_subtotals = []
    for i, cell, is_candidate in iter_rows(ws):
        if is_candidate:
            candidates.append(cell)
        else:
            non_candidate_subtotals.append(cell)
        subtotals.append(cell)

    results = {
        'candidates': candidates,
        'non_candidate_subtotals': non_candidate_subtotals,
        'subtotals': subtotals,
    }
    return results


def parse_workbook(path):
    """
    Parse a workbook, and return a dict of results.
    """
    results = {}
    wb = openpyxl.load_workbook(filename=path)
    _log.info(f'loaded workbook with sheets: {wb.sheetnames}')
    results = {}
    # TODO: parse the contest name from Sheet 1.
    sheet_2_results = parse_sheet_2(wb)
    results.update(sheet_2_results)
    return results
